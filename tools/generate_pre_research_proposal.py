from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables"
ASSET_DIR = OUTPUT_DIR / "assets"
OUTPUT_FILE = OUTPUT_DIR / "抽水蓄能电站建设期水土保持措施智能辅助决策系统研发_预研项目方案.docx"

COLOR_DARK = "24333B"
COLOR_GREEN = "0B6B57"
COLOR_GREEN_LIGHT = "DDEFEA"
COLOR_BLUE = "2F6690"
COLOR_BLUE_LIGHT = "E5EEF6"
COLOR_AMBER = "C67A16"
COLOR_AMBER_LIGHT = "F7EAD7"
COLOR_GRAY = "66737B"
COLOR_GRAY_LIGHT = "EEF1F2"
COLOR_WHITE = "FFFFFF"
COLOR_LINE = "B8C2C7"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = COLOR_LINE, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None,
                  align=WD_ALIGN_PARAGRAPH.LEFT, size: float = 9.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]],
              widths: Sequence[float] | None = None, header_fill: str = COLOR_GREEN) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        set_cell_text(hdr.cells[idx], header, bold=True, color=COLOR_WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[idx], header_fill)
        set_cell_border(hdr.cells[idx])
        set_cell_margins(hdr.cells[idx])
        hdr.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            hdr.cells[idx].width = Cm(widths[idx])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        if row_index % 2 == 1:
            for c in row.cells:
                set_cell_shading(c, "F7F9F9")
        for idx, value in enumerate(values):
            set_cell_text(row.cells[idx], str(value), size=9.2)
            set_cell_border(row.cells[idx])
            set_cell_margins(row.cells[idx])
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                row.cells[idx].width = Cm(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_para(doc: Document, text: str = "", *, bold_prefix: str | None = None,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent: bool = True,
             color: str | None = None, size: float = 10.5, space_after: float = 5.0) -> object:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        first.bold = True
        remainder = p.add_run(text[len(bold_prefix):])
        runs = (first, remainder)
    else:
        runs = (p.add_run(text),)
    for run in runs:
        run.font.name = "SimSun"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def add_bullet(doc: Document, text: str, *, level: int = 0, color: str | None = None) -> object:
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.55)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(10.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def add_numbered(doc: Document, text: str) -> object:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(10.5)
    return p


def add_callout(doc: Document, title: str, body: str, *, color: str = COLOR_GREEN,
                fill: str = COLOR_GREEN_LIGHT) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(0.25)
    table.columns[1].width = Cm(15.7)
    set_cell_shading(table.cell(0, 0), color)
    set_cell_border(table.cell(0, 0), color, 0)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_border(table.cell(0, 1), fill, 0)
    set_cell_margins(table.cell(0, 1), 140, 180, 140, 180)
    cell = table.cell(0, 1)
    cell.text = ""
    p1 = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(3)
    r1 = p1.add_run(title)
    r1.bold = True
    r1.font.name = "Microsoft YaHei"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = RGBColor.from_string(color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.35
    r2 = p2.add_run(body)
    r2.font.name = "SimSun"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    r2.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(5)

    heading_specs = {
        "Title": (25, COLOR_DARK, "Microsoft YaHei", True, 0),
        "Subtitle": (13, COLOR_GRAY, "Microsoft YaHei", False, 0),
        "Heading 1": (16, COLOR_GREEN, "Microsoft YaHei", True, 10),
        "Heading 2": (13.5, COLOR_DARK, "Microsoft YaHei", True, 7),
        "Heading 3": (11.5, COLOR_BLUE, "Microsoft YaHei", True, 5),
    }
    for name, (size, color, font, bold, before) in heading_specs.items():
        style = styles[name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(10.5)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    core = doc.core_properties
    core.title = "抽水蓄能电站建设期水土保持措施智能辅助决策系统研发——预研项目方案"
    core.subject = "课题五·研究内容3专项预研"
    core.author = "项目组"
    core.keywords = "抽水蓄能；水土保持；智能辅助决策；规则引擎；QGIS；C端"
    core.comments = "基于现有 QGIS/C++ 离线桌面端架构重构"


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("抽水蓄能电站建设期水土保持措施智能辅助决策系统研发")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(COLOR_GRAY)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), COLOR_GREEN)
    borders.append(bottom)
    pPr.append(borders)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fr = fp.add_run("专项预研方案  |  ")
    fr.font.name = "Microsoft YaHei"
    fr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(COLOR_GRAY)
    add_page_field(fp, "PAGE")


def font(size: int, *, bold: bool = False, color: str = COLOR_DARK):
    path = Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc")
    fallback = Path("C:/Windows/Fonts/simhei.ttf")
    return ImageFont.truetype(str(path if path.exists() else fallback), size)


def rounded(draw: ImageDraw.ImageDraw, box, radius: int, fill: str, outline: str | None = None, width: int = 2):
    draw.rounded_rectangle(box, radius=radius, fill=f"#{fill}", outline=f"#{outline}" if outline else None, width=width)


def wrapped_lines(draw: ImageDraw.ImageDraw, text: str, target_width: int, used_font) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if draw.textbbox((0, 0), candidate, font=used_font)[2] <= target_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw: ImageDraw.ImageDraw, box, text: str, used_font, fill: str, max_width: int | None = None,
                       line_gap: int = 8) -> None:
    x1, y1, x2, y2 = box
    lines = wrapped_lines(draw, text, max_width or (x2 - x1 - 20), used_font)
    heights = [draw.textbbox((0, 0), line, font=used_font)[3] for line in lines]
    total = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=used_font)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=used_font, fill=f"#{fill}")
        y += height + line_gap


def arrow(draw: ImageDraw.ImageDraw, start, end, color: str = COLOR_GRAY, width: int = 5) -> None:
    draw.line([start, end], fill=f"#{color}", width=width)
    x, y = end
    if abs(end[0] - start[0]) > abs(end[1] - start[1]):
        direction = 1 if end[0] > start[0] else -1
        points = [(x, y), (x - 14 * direction, y - 9), (x - 14 * direction, y + 9)]
    else:
        direction = 1 if end[1] > start[1] else -1
        points = [(x, y), (x - 9, y - 14 * direction), (x + 9, y - 14 * direction)]
    draw.polygon(points, fill=f"#{color}")


def create_architecture_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 1180), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = font(42, bold=True)
    section_font = font(27, bold=True)
    box_font = font(23, bold=True)
    small_font = font(19)
    draw.text((70, 45), "基于现有 C 端的目标产品架构", font=title_font, fill=f"#{COLOR_DARK}")
    draw.text((70, 105), "复用 QGIS 空间底座与端侧 AI，新增可独立测试、可版本化的水保决策域", font=small_font, fill=f"#{COLOR_GRAY}")

    layers = [
        (170, 350, "交互层", COLOR_DARK, ["工程工作台", "地图选取/识别", "引导式参数表单", "决策结果与人工复核"]),
        (380, 570, "应用编排层", COLOR_BLUE, ["扰动单元管理", "参数提取", "方案比选", "报告任务/审计"]),
        (600, 850, "水保决策域（本次核心新增）", COLOR_GREEN, ["规则引擎", "设计参数计算", "工程量与投资", "验收清单", "置信度/复核门"]),
        (880, 1080, "数据与适配层", COLOR_AMBER, ["QGIS/GDAL", "GeoPackage/SQLite", "JSON 规则包", "OpenXML 报告模板"]),
    ]
    for top, bottom, label, color, boxes in layers:
        draw.rounded_rectangle((60, top, 1740, bottom), radius=12, fill="#FAFBFB", outline=f"#{color}", width=3)
        draw.text((88, top + 20), label, font=section_font, fill=f"#{color}")
        available = 1560
        gap = 18
        box_width = int((available - gap * (len(boxes) - 1)) / len(boxes))
        x = 150
        box_top = top + 75
        box_bottom = bottom - 28
        for index, item in enumerate(boxes):
            fill = COLOR_GREEN_LIGHT if top == 565 else (COLOR_BLUE_LIGHT if top == 340 else (COLOR_AMBER_LIGHT if top == 830 else COLOR_GRAY_LIGHT))
            rounded(draw, (x, box_top, x + box_width, box_bottom), 8, fill, color, 2)
            draw_centered_text(draw, (x + 10, box_top + 5, x + box_width - 10, box_bottom - 5), item, box_font, COLOR_DARK)
            x += box_width + gap

    draw.text((85, 1110), "现有可复用：QGIS 工程/图层树、二维三维地图、ONNX 扰动识别、SAM2 分割、人工核查、离线工程目录", font=small_font, fill=f"#{COLOR_GRAY}")
    image.save(path, quality=95)


def create_flow_diagram(path: Path) -> None:
    image = Image.new("RGB", (1800, 850), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = font(42, bold=True)
    box_font = font(23, bold=True)
    small_font = font(18)
    draw.text((70, 42), "技术路线与闭环工作流", font=title_font, fill=f"#{COLOR_DARK}")
    draw.text((70, 102), "从空间扰动单元出发，保证每个推荐结果可计算、可追溯、可复核、可导出", font=small_font, fill=f"#{COLOR_GRAY}")

    steps = [
        ("1", "知识源数字化", "标准/案例/专家经验", COLOR_BLUE),
        ("2", "三库与计算库", "规则、单价、验收、参数公式", COLOR_GREEN),
        ("3", "空间参数获取", "识别图斑/人工绘制/图层提取", COLOR_AMBER),
        ("4", "智能决策", "匹配、计算、冲突检查、复核门", COLOR_GREEN),
        ("5", "三件套输出", "措施建议、投资估算、验收清单", COLOR_BLUE),
        ("6", "试点迭代", "对比评审、误差分析、版本固化", COLOR_AMBER),
    ]
    y1, y2 = 235, 545
    gap = 30
    width = 250
    x = 55
    for index, (num, title, desc, color) in enumerate(steps):
        rounded(draw, (x, y1, x + width, y2), 12, "FFFFFF", color, 4)
        rounded(draw, (x + 84, y1 + 24, x + 166, y1 + 106), 41, color)
        draw_centered_text(draw, (x + 84, y1 + 24, x + 166, y1 + 106), num, font(34, bold=True), COLOR_WHITE)
        draw_centered_text(draw, (x + 18, y1 + 120, x + width - 18, y1 + 190), title, box_font, COLOR_DARK)
        draw_centered_text(draw, (x + 20, y1 + 196, x + width - 20, y2 - 20), desc, small_font, COLOR_GRAY)
        if index < len(steps) - 1:
            arrow(draw, (x + width + 6, (y1 + y2) // 2), (x + width + gap - 6, (y1 + y2) // 2), COLOR_GRAY, 4)
        x += width + gap

    rounded(draw, (360, 650, 1440, 785), 10, COLOR_GREEN_LIGHT, COLOR_GREEN, 3)
    draw_centered_text(draw, (390, 665, 1410, 715), "形成企业级可积累资产", font(25, bold=True), COLOR_GREEN)
    draw_centered_text(draw, (390, 720, 1410, 772), "规则版本 + 来源条文 + 适用条件 + 计算过程 + 专家复核记录 + 试点证据", small_font, COLOR_DARK)
    arrow(draw, (1540, 545), (1445, 695), COLOR_AMBER, 4)
    image.save(path, quality=95)


def add_picture(doc: Document, path: Path, caption: str, width_cm: float = 16.1) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(7)
    run = cp.add_run(caption)
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(COLOR_GRAY)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("课题五 · 研究内容3专项预研")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(COLOR_GREEN)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(28)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.25
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run("抽水蓄能电站建设期\n水土保持措施智能辅助决策系统研发")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(COLOR_DARK)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(28)
    sr = sub.add_run("预研项目方案（V1.0）")
    sr.font.name = "Microsoft YaHei"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sr.font.size = Pt(18)
    sr.font.bold = True
    sr.font.color.rgb = RGBColor.from_string(COLOR_BLUE)

    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.cell(0, 0)
    set_cell_shading(cell, COLOR_GREEN_LIGHT)
    set_cell_border(cell, COLOR_GREEN, 8)
    set_cell_margins(cell, 230, 260, 230, 260)
    cell.text = ""
    bp = cell.paragraphs[0]
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = bp.add_run("经费规模约 10 万元  ·  实施周期 12 个月  ·  核心成果为可推广转化产品")
    br.font.name = "Microsoft YaHei"
    br._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    br.font.size = Pt(12)
    br.font.bold = True
    br.font.color.rgb = RGBColor.from_string(COLOR_GREEN)

    for _ in range(5):
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    entries = [
        ("承担单位", "____________________________"),
        ("项目负责人", "____________________________"),
        ("编制日期", "2026 年 8 月"),
        ("版本说明", "基于现有 QGIS/C++ C 端架构梳理"),
    ]
    for row, (label, value) in zip(meta.rows, entries):
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(9)
        set_cell_text(row.cells[0], label, bold=True, color=COLOR_DARK, align=WD_ALIGN_PARAGRAPH.RIGHT, size=10.5)
        set_cell_text(row.cells[1], value, color=COLOR_GRAY, align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5)
        for c in row.cells:
            set_cell_border(c, COLOR_WHITE, 0)
            set_cell_margins(c, 100, 180, 100, 180)

    doc.add_page_break()


def add_contents(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    items = [
        "一、项目摘要",
        "二、现有 C 端架构基础与本次建设边界",
        "三、项目定位与立项逻辑",
        "四、研究目标与考核指标",
        "五、研究范围与对象",
        "六、研究任务与主要内容",
        "七、系统总体架构与数据架构",
        "八、决策模型与关键业务流程",
        "九、技术路线、关键难点与创新点",
        "十、试点验证与评价方法",
        "十一、现场应用价值",
        "十二、成果体系与推广转化",
        "十三、实施计划与组织保障",
        "十四、经费概算",
        "十五、风险分析与对策",
        "十六、与后续课题的衔接",
        "附录 A：V1.0 功能优先级清单",
        "附录 B：规则数据结构示例",
        "附录 C：主要依据与版本管理原则",
    ]
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor.from_string(COLOR_DARK)
    add_callout(doc, "阅读提示", "本方案中的“本项目”均指课题五研究内容3的专项预研；项目内部按五项研究任务组织，不再使用“研究内容1/2/3”的编号，避免层级混淆。", color=COLOR_BLUE, fill=COLOR_BLUE_LIGHT)
    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    configure_document(doc)
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc)
    add_contents(doc)

    doc.add_heading("一、项目摘要", level=1)
    add_para(doc, "本预研项目面向抽水蓄能电站建设期弃渣、挖填方、道路及施工场地等高频扰动场景，研发“水土保持措施智能辅助决策系统 V1.0”。系统以当前已形成的 QGIS/C++ 离线桌面端为载体，复用工程管理、影像和矢量数据、二维/三维地图、ONNX 扰动识别、SAM2 智能分割、人工核查与成果归档能力，在其上新增扰动单元建模、空间参数提取、可追溯规则引擎、设计参数初算、投资估算、验收清单与 Word/Excel 报告输出。")
    add_para(doc, "V1.0 不另起一套独立 B/S 系统，而是采用“桌面 GIS 主端 + 本地规则与计算组件 + 标准化数据接口”的增量路线。该路线能够直接承接现有扰动识别图斑，将“发现扰动”延伸到“提出措施、估算投资、指导验收”，同时保留后续封装服务 API、接入课题五大模型问答系统和数字孪生平台的能力。")
    add_callout(doc, "方案核心判断", "10 万元预研经费应集中形成可验证的决策内核与可运行产品闭环，不宜同时建设完整 Web 平台、知识图谱、大模型问答和全专业设计系统。V1.0 的核心资产是规则库及其计算、追溯、验证机制。")
    add_table(doc, ["项目要素", "建议口径"], [
        ["项目属性", "课题五研究内容3专项预研；以产品化原型和企业规则资产为主"],
        ["建设周期", "12 个月"],
        ["经费规模", "约 10 万元"],
        ["产品形态", "Windows 离线桌面端模块 + 三库数据包 + 报告模板 + 使用/部署文档"],
        ["V1.0 深度范围", "渣场、边坡、施工道路、施工生产生活区 4 类典型单元形成完整决策闭环"],
        ["扩展范围", "统一数据模型覆盖 8～10 类扰动单元，料场、表土堆场等通过模板继续扩展"],
        ["核心输出", "措施建议清单、投资估算表、验收核查表及一份 Word 建议书"],
        ["试点方式", "1 个在建抽蓄电站、3～5 个典型扰动单元、2～3 名专家对比评审"],
    ], widths=[3.3, 12.7])

    doc.add_heading("二、现有 C 端架构基础与本次建设边界", level=1)
    doc.add_heading("（一）现有架构与能力基础", level=2)
    add_para(doc, "经对当前代码仓和运行形态梳理，现有项目是基于 QGIS 源码进行 C++/Qt 原生扩展的 Windows 桌面 GIS 产品，不是普通浏览器前端。业务能力由原生工作台控制器接入 QGIS 工程、图层树、地图画布、矢量编辑和数据提供器，推理侧使用 ONNX Runtime，已集成施工扰动分割模型和 SAM2 交互式精细分割模型，并采用本地工程目录完成离线归档。该基础与建设现场弱网、空间数据量大、成果需落图核查的特点高度匹配。")
    add_table(doc, ["现有能力", "当前实现基础", "本项目复用方式"], [
        ["工程与目录", "QGIS .qgz 工程；imagery、vectors、models、results、screenshots、reports、cache 标准目录", "扩展为抽蓄水保工程模板，继续按项目离线归档"],
        ["空间数据", "GeoTIFF、SHP/GPKG/GeoJSON 导入；QGIS/GDAL 坐标转换、渲染和空间分析", "承接 DEM、遥感影像、扰动单元、汇水区、措施布置等数据"],
        ["智能识别", "ONNX 扰动图斑识别、分块推理、跨瓦片抑制与地理坐标回写", "识别图斑可直接转为待决策扰动单元"],
        ["智能分割", "SAM2 点选/框选提示分割，输出可编辑面图斑", "用于渣场、边坡、场地边界快速精细化"],
        ["人工核查", "矢量新增、节点编辑、选择、删除、状态记录与结果表格", "用于补录决策参数、修正边界、确认措施建议"],
        ["地图与展示", "二维/三维视图、图层树、业务分组、定位与显隐", "在地图上联动查看扰动单元、推荐措施和风险提示"],
        ["成果输出", "批量截图与 CSV 索引，工程 reports 目录已预留", "新增 Word/Excel/JSON 模板化成果输出"],
        ["部署方式", "C++ 原生应用、ONNX 模型随应用本地部署", "保持断网可用；规则包和单价包可离线升级"],
    ], widths=[2.6, 6.5, 6.9])

    doc.add_heading("（二）本次新增能力与工程边界", level=2)
    add_table(doc, ["层面", "本次新增", "明确不做/后续做"], [
        ["业务", "抽蓄工程信息、扰动单元台账、决策任务、方案比选、人工复核", "不建设覆盖所有行业的通用水保设计平台"],
        ["数据", "扰动单元参数模型、规则库、单价库、验收指标库、计算记录", "不一次性完成全国全区域实时价格库"],
        ["算法", "规则匹配、参数初算、工程量估算、投资汇总、验收映射", "不替代拦挡结构稳定计算、地勘、施工图和法定审查"],
        ["产品", "现有 C 端内新增“水保决策”工作空间和报告组件", "本阶段不另建完整 B/S 门户、知识图谱和大模型问答层"],
        ["接口", "统一 JSON 输入/输出、规则包导入导出、未来服务接口 DTO", "V1.0 不强制部署服务器或建设复杂微服务"],
    ], widths=[2.3, 6.8, 6.9])
    add_callout(doc, "工程实施建议", "现有生态修复控制器已承担较多界面与业务职责。新增水保决策功能应作为独立业务模块落位，核心规则、计算、估算和报告逻辑采用无界面的领域服务，避免继续堆入单一控制器，便于单元测试和后续服务化。", color=COLOR_AMBER, fill=COLOR_AMBER_LIGHT)

    doc.add_heading("三、项目定位与立项逻辑", level=1)
    doc.add_heading("（一）项目定位", level=2)
    add_para(doc, "本项目是课题五“大模型合规问答系统”的先行切入点，也是现有生态环境遥感智能核查 C 端从“扰动发现与人工核查”向“措施决策与闭环管理”延伸的核心模块。产品面向建设单位环水保管理人员、施工单位技术人员、水保方案编制/监理人员和验收自查人员，提供方案阶段和施工调整阶段的快速辅助决策。")
    doc.add_heading("（二）优先研发辅助决策的理由", level=2)
    add_numbered(doc, "问答系统的基础壁垒主要在知识组织，辅助决策的核心壁垒在规则模型。把标准条文、案例和专家经验转化为可计算、可追溯、可验证的规则，是最难复制且可持续积累的企业技术资产。")
    add_numbered(doc, "10 万元无法覆盖完整的大模型系统和全专业设计能力，但可以把“扰动单元输入—措施推荐—投资初算—验收清单—报告输出”做成一个可试用、可验收的小闭环。")
    add_numbered(doc, "当前 C 端已经具备空间数据、识别、分割、编辑和离线工程能力，新增决策域即可形成端到端产品，避免重复投入地图、数据和部署底座。")
    add_numbered(doc, "决策内核可先作为桌面模块在单站试用，后续通过标准 JSON/服务封装接入大模型 RAG、数字孪生或企业 Web 平台，前期投资不重复。")

    doc.add_heading("四、研究目标与考核指标", level=1)
    doc.add_heading("（一）总体目标", level=2)
    add_para(doc, "研发“抽水蓄能电站建设期水土保持措施智能辅助决策系统 V1.0”。用户可从地图选择识别/绘制的扰动单元，由系统自动提取面积、坡度等可计算参数，补充土壤、降雨、防治分区等业务参数后，自动生成措施组合、关键设计参数初算、工程量和投资估算、验收要点，并保留规则来源、计算过程、缺失参数和人工复核记录。")
    doc.add_heading("（二）分目标", level=2)
    add_bullet(doc, "形成统一的扰动单元数据模型和 8～10 类分类框架，优先完成 4 类典型单元的端到端闭环。")
    add_bullet(doc, "建立可版本化、可追溯的决策规则库、综合单价库、验收指标库及关键参数计算库。")
    add_bullet(doc, "在现有 QGIS/C++ 客户端中形成可离线运行的引导式决策工作台，支持地图与表单双入口。")
    add_bullet(doc, "在 1 个在建抽蓄电站完成 3～5 个典型单元试点，对规则适用性、投资偏差和报告可用性进行验证。")
    add_bullet(doc, "完成 V1.0 产品封装、使用手册、部署包和知识产权材料，具备复制到其他站点的条件。")

    doc.add_heading("（三）建议考核指标", level=2)
    add_table(doc, ["指标类别", "V1.0 建议指标", "验收方法"], [
        ["场景覆盖", "建立 8～10 类统一分类；渣场、边坡、施工道路、施工生产生活区 4 类形成完整闭环", "功能清单、测试用例、试点记录"],
        ["规则资产", "不少于 60 条可执行规则/参数约束；发布规则 100% 绑定来源、版本和适用条件", "规则库审查与自动校验报告"],
        ["单价资产", "收录不少于 80 项常用措施，目标 80～120 项；支持地区、基期和调整系数", "单价库清单与抽查"],
        ["验收映射", "V1.0 措施目录均建立验收指标、核查方式、阈值或判定依据映射", "覆盖率检查"],
        ["功能闭环", "参数录入、规则匹配、计算、估算、验收、Word/Excel/JSON 输出均可离线运行", "现场演示与断网测试"],
        ["响应性能", "资料完整时单个扰动单元规则计算与结果生成目标 ≤3 秒，不含大范围 GIS 预处理", "指定基准机计时"],
        ["措施合理性", "试点专家一致认定的必要措施无原则性漏项；一级措施组合与专家共识匹配率目标 ≥80%", "盲评表和差异分析"],
        ["估算精度", "口径、价格基期和输入工程量一致条件下，试点单元汇总投资相对偏差目标 ≤15%", "与批复/复核结果对比"],
        ["可追溯性", "每次结果保存输入快照、规则版本、命中规则、公式与人工修改记录", "结果包和日志抽查"],
        ["产品成果", "V1.0 安装/离线部署包 1 套、软著至少 1 项、试点应用报告 1 份", "成果文件与受理材料"],
    ], widths=[2.5, 8.2, 5.3])
    add_callout(doc, "指标适用条件", "“投资偏差≤15%”仅适用于同一计价口径、价格基期一致、基础资料完整的试点样本，不承诺代替正式概（估）算。对于高风险、资料缺失或超出规则适用域的工况，系统必须输出人工复核提示而非强行给出确定结论。", color=COLOR_AMBER, fill=COLOR_AMBER_LIGHT)

    doc.add_heading("五、研究范围与对象", level=1)
    doc.add_heading("（一）扰动单元分类框架", level=2)
    add_table(doc, ["分类层级", "扰动单元", "V1.0 处理深度"], [
        ["P0 核心", "渣场（沟道型/坡地型）", "完整规则、参数计算、投资和验收闭环；重点试点"],
        ["P0 核心", "边坡（土质挖方/填方/岩质）", "完整规则、措施比选与典型工程量闭环"],
        ["P0 核心", "施工道路（永久/临时、挖填段）", "完整排水、防护、绿化及临时措施闭环"],
        ["P0 核心", "施工生产生活区", "完整场地排水、沉沙、临时苫盖、恢复闭环"],
        ["P1 扩展", "料场、表土堆场", "建立数据模板和基础措施规则，视资料完成扩展验证"],
        ["P1 扩展", "上下库连接工程、输水发电系统施工区", "纳入分类与接口，复杂结构由专业复核"],
        ["P1 扩展", "施工营地、临时堆土及其他扰动区", "通用兜底规则和人工复核机制"],
    ], widths=[2.5, 5.5, 8.0])
    doc.add_heading("（二）输入、处理与输出边界", level=2)
    add_table(doc, ["环节", "主要内容"], [
        ["空间输入", "扰动图斑、DEM/等高线、影像、工程分区、汇水边界、土壤/植被/降雨分区图层（可选）"],
        ["业务输入", "单元类型、阶段、坡率/堆高/堆渣量、土壤类型、降雨、防治分区、设计频率、施工时段、价格基期等"],
        ["自动处理", "面积/长度/平均坡度提取、规则匹配、关键参数初算、工程量推导、单价套用、验收映射、完整性/冲突检查"],
        ["人工补充", "无法从空间图层可靠获取的地质、材料、工艺、设计标准和本地价格参数"],
        ["输出成果", "措施建议书（Word）、投资估算表（Excel）、验收核查表（Excel）、决策结果包（JSON/GeoPackage）"],
        ["责任边界", "系统用于方案比选、变更研判和自查；重大拦挡工程、结构稳定及施工图必须由有资质专业人员复核"],
    ], widths=[3.0, 13.0])

    doc.add_heading("六、研究任务与主要内容", level=1)
    add_para(doc, "本专项预研在课题体系中属于“研究内容3”，项目内部按以下五项研究任务组织。五项任务以规则资产为主线、以 C 端产品为载体、以试点证据为验收依据。")

    doc.add_heading("任务一：决策知识源数字化与规则体系构建（核心）", level=2)
    doc.add_heading("1. 知识源梳理", level=3)
    add_bullet(doc, "梳理生产建设项目水土保持技术、设计、防治、监测、质量评定和设施验收等现行标准及水利部、流域机构、国网公司相关制度文件。")
    add_bullet(doc, "收集 5～10 个已建/在建抽水蓄能电站水土保持方案、变更、监理、监测和验收材料；对敏感信息进行脱敏，形成案例样本表。")
    add_bullet(doc, "访谈 3～5 名方案编制、设计、监理或验收专家，提炼标准未直接给出的工程适用条件、禁配条件、优先级和复核门。")
    doc.add_heading("2. 规则数据模型", level=3)
    add_para(doc, "建立“扰动单元类型 × 立地/工程条件 × 建设阶段 → 措施组合 + 参数约束 + 复核要求”的规则框架。每条规则至少包含规则编号、条件表达式、输出措施、优先级、互斥/联动关系、适用区域、标准版本、来源条款、专家确认状态、置信等级和生效日期。规则表达式采用白名单字段与安全运算符，不执行任意脚本。")
    doc.add_heading("3. 规则治理", level=3)
    add_bullet(doc, "建立草稿、评审、发布、停用四态生命周期；发布规则只读，修改通过新版本完成。")
    add_bullet(doc, "建设冲突、重复、空条件、死规则和缺失来源自动校验；所有决策保存命中链路。")
    add_bullet(doc, "设置兜底推荐与人工复核门。低覆盖度、关键参数缺失、超出范围或高风险工程只输出原则性建议和缺项清单。")

    doc.add_heading("任务二：空间参数提取与关键参数计算", level=2)
    doc.add_heading("1. 扰动单元建模与空间参数提取", level=3)
    add_para(doc, "在现有识别/人工核查图斑基础上建立“扰动单元”主对象，关联工程分区、影像期次、核查状态和决策版本。利用 QGIS/GDAL 能力自动计算面积、周长、坡度统计、坡向、相对高差等；具备可靠基础数据时进一步计算汇水面积和与沟道/道路的空间关系。所有自动值标记来源图层、时间和算法，允许人工覆盖但保留原值。")
    doc.add_heading("2. 关键参数初算", level=3)
    add_table(doc, ["计算模块", "输入", "输出", "V1.0 边界"], [
        ["截排水水力初算", "汇水面积、设计暴雨/强度、径流系数、纵坡、糙率", "设计流量、推荐断面、流速及校核提示", "用于方案初选；地区暴雨公式和安全超高可配置"],
        ["拦挡规模估算", "堆渣量、堆高、沟谷条件、下游敏感性", "工程级别建议、控制尺寸区间、专业复核项", "不替代坝体稳定、地基和结构设计"],
        ["边坡防护选型", "坡高、坡率、土/岩性、坡面径流、施工阶段", "工程/生态/临时防护组合及参数范围", "超高边坡、滑坡等自动触发专项设计"],
        ["植物措施推荐", "海拔、气候、土壤、坡向、恢复目标、管护条件", "草+灌+藤/乔灌草组合、播种/栽植量及养护要点", "V1.0 先建立试点区域适生植物表"],
    ], widths=[3.0, 4.2, 4.5, 4.3])
    add_para(doc, "所有公式应保存公式编号、输入值、单位、结果、适用范围和校核状态。统一采用单位字典和量纲校验，防止平方米/公顷、毫米/米、自然坡度/坡率等常见口径错误。")

    doc.add_heading("任务三：投资估算与验收要点生成", level=2)
    doc.add_heading("1. 投资估算引擎", level=3)
    add_bullet(doc, "建立不少于 80 项常用措施综合单价条目，记录措施编码、单位、基价、人工/材料/机械构成、地区、价格基期、依据版本和调整系数。")
    add_bullet(doc, "形成“措施—工程量公式—单价—费用分类”映射，自动汇总工程措施费、植物措施费、临时措施费，并按配置计算独立费用、预备费等。")
    add_bullet(doc, "支持项目级人工费、材料价、地区和基期调整；输出原始基价、调整项和结果，避免只给总价无法解释。")
    add_bullet(doc, "提供方案 A/B 对比，展示投资差额、适用条件、维护要求和风险提示，不以最低价作为唯一推荐依据。")
    doc.add_heading("2. 验收要点生成", level=3)
    add_bullet(doc, "建立“措施—验收指标—核查方式—阈值/判定依据—资料要求”映射库，覆盖 V1.0 措施目录。")
    add_bullet(doc, "按单位工程、分部工程、单元工程层次组织核查项，支持现场量测、无人机/影像核查、资料核查和照片留痕。")
    add_bullet(doc, "服务现行六项防治指标自评：水土流失治理度、土壤流失控制比、渣土防护率、表土保护率、林草植被恢复率、林草覆盖率。")
    add_bullet(doc, "输出 Excel 自查表，包含责任人、计划/完成日期、问题、整改状态和证据链接字段，便于日常闭环。")

    doc.add_heading("任务四：C 端原型系统开发与产品化封装", level=2)
    doc.add_heading("1. 产品工作流", level=3)
    add_para(doc, "在当前桌面端新增“水保决策”工作空间，推荐交互顺序为：选择/创建扰动单元 → 自动提取空间参数 → 引导补齐业务参数 → 运行规则与计算 → 查看命中依据和风险 → 方案比选/人工确认 → 一键生成成果。地图与表单保持双向联动，用户可从图斑进入决策，也可从台账定位图斑。")
    doc.add_heading("2. 软件模块", level=3)
    add_table(doc, ["模块", "主要功能", "建议技术落位"], [
        ["水保决策工作台", "单元台账、分步表单、结果树、方案比选、复核", "Qt Widgets/QGIS Dock；独立控制器"],
        ["扰动单元服务", "图斑关联、参数快照、状态流转、地图联动", "QGIS API + GeoPackage 仓储"],
        ["规则引擎", "条件解析、优先级、冲突、命中链、覆盖度与复核门", "纯 C++ 领域服务；JSON/SQLite 规则仓储"],
        ["计算与估算", "水力、工程量、费用分类、调整系数、方案对比", "纯 C++ 计算库；单元测试覆盖"],
        ["验收生成", "措施映射、层级清单、六项指标关联", "本地数据仓储 + 模板服务"],
        ["报告组件", "DOCX、XLSX、JSON/GeoPackage 输出", "模板化 OpenXML 适配器，输出至 reports"],
        ["规则包管理", "导入、校验、版本查看、启停、备份", "V1.0 提供受控导入；可视化编辑器列为 P1"],
    ], widths=[3.0, 6.0, 7.0])
    doc.add_heading("3. 离线部署与接口预留", level=3)
    add_para(doc, "产品继续采用 Windows 单机离线部署，模型、规则、单价、验收库和模板随安装包交付。对外以版本化 JSON DTO 作为稳定边界，输入扰动单元与参数，输出措施、计算、投资、验收和追溯信息；后续可在不重写核心逻辑的前提下封装为 REST 服务。")

    doc.add_heading("任务五：试点验证、迭代与成果固化", level=2)
    add_bullet(doc, "选取 1 个在建抽水蓄能电站，优先覆盖 1 个沟道型渣场、2 段典型边坡、1 条施工道路和 1 个施工生产生活区中的 3～5 个单元。")
    add_bullet(doc, "对比正式水保方案/变更成果，采用统一输入口径运行系统；邀请 2～3 名未参与规则编制的专家进行独立评审。")
    add_bullet(doc, "形成差异分类：输入不足、规则缺失、规则冲突、计算误差、单价口径差异、专家偏好差异；逐项闭环修订。")
    add_bullet(doc, "冻结 V1.0 规则包、单价包、验收包、报告模板和软件版本，出具试点应用报告、测试报告和发布说明。")

    doc.add_heading("七、系统总体架构与数据架构", level=1)
    architecture_path = ASSET_DIR / "target_architecture.png"
    add_picture(doc, architecture_path, "图 1  基于现有 C 端的目标产品架构")
    doc.add_heading("（一）架构原则", level=2)
    add_bullet(doc, "空间底座复用：继续使用 QGIS/GDAL 的工程、渲染、坐标、几何和分析能力。")
    add_bullet(doc, "决策域解耦：规则、公式、估算和验收不依赖界面对象，可独立测试、批处理和服务化。")
    add_bullet(doc, "数据本地优先：弱网/断网可完成核心流程；规则包可通过受控文件离线升级。")
    add_bullet(doc, "结果可追溯：输入、版本、命中规则、计算、人工修改和输出文件相互关联。")
    add_bullet(doc, "人机协同：空间 AI 的“模型置信度”与规则决策的“覆盖度/确定性”分开表达，低确定性结果必须人工复核。")

    doc.add_heading("（二）数据分层与存储建议", level=2)
    add_table(doc, ["数据层", "建议载体", "主要内容", "版本策略"], [
        ["工程空间数据", "项目 GeoPackage + QGIS 工程", "扰动单元、措施布置、工程分区、核查记录、空间关系", "按工程保存，关键成果可创建快照"],
        ["决策知识数据", "knowledge.db（SQLite）", "规则、措施目录、参数字典、公式、来源、版本和评审记录", "只读发布包；版本号+校验值"],
        ["交换规则包", "JSON", "规则及依赖项的导入导出、审查和后续 API 数据", "语义版本；保留向后兼容迁移"],
        ["单价数据", "SQLite/Excel 导入模板", "地区、基期、综合单价、构成、调整系数", "与规则包解耦，支持年度更新"],
        ["报告模板", "DOCX/XLSX 模板", "建议书、估算表、核查表版式与字段映射", "模板 ID + 版本号"],
        ["运行与审计", "项目 SQLite/JSON 日志", "决策任务、输入快照、命中链、人工修订、导出清单", "只追加关键记录，可追溯"],
    ], widths=[2.5, 3.3, 6.6, 3.6])

    doc.add_heading("（三）与现有工程目录的衔接", level=2)
    add_table(doc, ["现有目录", "扩展后的用途"], [
        ["imagery", "遥感影像、正射影像、必要时的历史影像"],
        ["vectors", "工程分区、道路、沟道、扰动单元来源矢量"],
        ["models", "ONNX 扰动识别与交互分割模型"],
        ["results", "扰动单元、措施建议空间成果、决策结果数据库"],
        ["screenshots", "单元定位图、措施示意与核查照片索引"],
        ["reports", "Word 建议书、Excel 估算表/核查表、JSON 结果包"],
        ["cache", "DEM 派生坡度、汇水分析、临时渲染和模型推理缓存"],
    ], widths=[3.2, 12.8])

    doc.add_heading("八、决策模型与关键业务流程", level=1)
    doc.add_heading("（一）决策输入模型", level=2)
    add_table(doc, ["参数组", "代表参数", "获取方式"], [
        ["身份与阶段", "工程、分区、单元类型/子类型、施工阶段、影像期次", "工程属性/人工选择"],
        ["几何与地形", "面积、长度、坡高、坡率、坡向、相对高差、集水面积", "GIS 自动提取+人工校核"],
        ["工程规模", "堆渣量、最大堆高、道路等级/宽度、边坡结构", "设计资料/人工录入"],
        ["立地条件", "土壤/岩性、侵蚀强度、植被、海拔、防治分区", "专题图层/人工录入"],
        ["水文气象", "多年平均降雨、设计暴雨、径流系数、敏感目标", "参数表/图层/人工录入"],
        ["计价条件", "地区、价格基期、人工/材料调整、费率", "项目级配置"],
        ["质量与验收", "目标标准、核查方式、阈值、资料完整性", "验收库+项目配置"],
    ], widths=[2.5, 8.0, 5.5])

    doc.add_heading("（二）规则执行逻辑", level=2)
    add_numbered(doc, "完整性检查：区分必填、条件必填和可选参数；缺少关键参数时先输出缺项，不进入强结论。")
    add_numbered(doc, "候选规则筛选：按单元类型、阶段、区域、标准版本和有效期形成候选集。")
    add_numbered(doc, "条件求值与优先级：计算规则条件，处理互斥、依赖和高风险优先规则。")
    add_numbered(doc, "措施组合：形成工程、植物、临时三类措施；合并重复项并校验必要措施和禁配关系。")
    add_numbered(doc, "参数与工程量：调用公式库生成关键尺寸初值、工程量、单位和适用范围提示。")
    add_numbered(doc, "投资与验收：套用价格版本，生成分类汇总；映射验收指标和核查方式。")
    add_numbered(doc, "结果分级：输出规则覆盖度、数据完整度和复核级别，保存命中链与人工确认记录。")

    doc.add_heading("（三）典型规则链示例", level=2)
    add_callout(doc, "示例：沟道型渣场", "当“单元类型=渣场、子类型=沟道型、堆渣量>50 万 m³”时，候选措施为拦挡工程、上游/周边截水、排洪系统、堆体分级及马道排水、渣面整治与植被恢复。系统继续根据集水面积、设计频率、堆高、下游敏感性和地质资料计算或筛选参数；若拦挡等级高、地质资料缺失或超出规则适用范围，则强制标注“专项设计/稳定复核”，不直接给出可施工结构尺寸。", color=COLOR_BLUE, fill=COLOR_BLUE_LIGHT)
    add_para(doc, "上述规则链与原始“直接给出浆砌石/混凝土坝高”的表达相比，更符合辅助决策产品边界：系统可以确定措施体系、给出参数区间和计算依据，但不会在缺少地勘、稳定和结构验算时输出貌似精确的工程结论。")

    doc.add_heading("（四）结果置信与复核分级", level=2)
    add_table(doc, ["等级", "判定特征", "系统行为"], [
        ["A 可直接参考", "必填参数完整；规则覆盖充分；无冲突；处于公式适用域", "输出推荐方案、计算、投资与验收清单，提示常规专业复核"],
        ["B 建议复核", "存在可补充参数、多个近似方案或地区经验差异", "给出 2～3 个备选并列明差异、缺项和专家确认点"],
        ["C 强制专业介入", "高风险工程、关键地质/水文资料缺失、规则冲突或超出适用域", "仅输出原则性措施、资料清单和专项设计要求，不给出确定结构结论"],
    ], widths=[2.6, 7.0, 6.4])

    doc.add_heading("九、技术路线、关键难点与创新点", level=1)
    flow_path = ASSET_DIR / "technical_flow.png"
    add_picture(doc, flow_path, "图 2  技术路线与闭环工作流")
    doc.add_heading("（一）关键技术难点及对策", level=2)
    add_table(doc, ["难点", "风险表现", "对策"], [
        ["标准条文规则化存在歧义", "同一条文在不同工程条件下理解不同", "规则与来源双绑定；记录解释说明；3～5 名专家交叉评审；发布前测试"],
        ["规则覆盖不全", "长尾工况被错误套用", "适用域、兜底规则、覆盖度和强制复核门；禁止无依据外推"],
        ["空间参数质量不一致", "坐标、分辨率、DEM 精度导致误差", "数据质量检查、来源标记、自动值可人工覆盖、重要参数二次确认"],
        ["规则冲突与版本漂移", "标准更新后新旧规则同时生效", "语义版本、有效期、优先级、冲突校验、决策输入与规则快照"],
        ["单价区域与基期差异", "同一措施价格差异大", "基价、地区、基期和调整系数分离；允许项目价覆盖且留痕"],
        ["高风险结果被误用", "辅助建议被当作施工图", "显著责任边界、结果分级、强制复核、报告中自动附适用条件和限制"],
        ["现有代码耦合", "新增逻辑进入大控制器后难测试", "决策域模块化；UI 仅做编排；核心服务纯 C++、独立测试"],
    ], widths=[3.0, 5.2, 7.8])

    doc.add_heading("（二）预期创新点", level=2)
    add_numbered(doc, "空间识别与规则决策贯通。把现有 ONNX/SAM2 扰动图斑从“结果展示”升级为具有参数、状态和决策版本的业务对象，实现从遥感发现到措施建议的闭环。")
    add_numbered(doc, "规则—出处—计算—验收一体化追溯。每项建议不仅回答“做什么”，还回答“为什么、怎么算、如何验”，形成可审计企业规则资产。")
    add_numbered(doc, "GIS 自动提参和专家规则协同。面积、坡度、空间关系等由 GIS 计算，地质、阶段和管理条件由人员补充，降低录入负担并避免黑箱决策。")
    add_numbered(doc, "面向风险的人机协同机制。将模型置信度、数据完整度、规则覆盖度和专业复核级别分开管理，对高风险/长尾工况主动收敛结论。")
    add_numbered(doc, "端侧离线内核与服务化接口并存。V1.0 在现场离线可用，同时通过统一 DTO 为后续大模型工具调用和数字孪生接入预留稳定边界。")

    doc.add_heading("十、试点验证与评价方法", level=1)
    doc.add_heading("（一）试点对象与样本", level=2)
    add_para(doc, "选取资料相对完整、建设活动正在开展且具备现场核查条件的 1 个抽水蓄能电站。样本优先选择沟道型渣场、典型挖/填方边坡、施工道路和施工生产生活区，最终确定 3～5 个单元。每个单元建立统一“金标准输入表”，确保系统与正式方案采用相同边界、参数和价格基期。")
    doc.add_heading("（二）验证步骤", level=2)
    add_numbered(doc, "资料准备：收集批复方案、变更、设计、监理/监测、计价资料和现场影像，完成脱敏与输入冻结。")
    add_numbered(doc, "系统独立运行：由未参与规则编制的操作人员按手册完成图斑确认、参数补录和结果输出。")
    add_numbered(doc, "专家盲评：隐藏正式方案结论，专家分别评价措施完整性、适用性、参数合理性、风险提示和报告可用性。")
    add_numbered(doc, "定量对比：计算措施匹配、原则性漏项、工程量差异、投资汇总偏差、规则覆盖度和操作耗时。")
    add_numbered(doc, "问题闭环：按缺数据、缺规则、错规则、公式/单位、价格口径、交互问题分类整改并回归测试。")
    add_numbered(doc, "版本冻结：发布软件、规则包、单价包、验收包和模板版本，形成可重复验证的发布基线。")

    doc.add_heading("（三）评价矩阵", level=2)
    add_table(doc, ["评价维度", "核心问题", "证据"], [
        ["正确性", "必要措施是否遗漏，禁配措施是否误选，参数是否超适用域", "专家评分、命中规则链、差异清单"],
        ["可解释性", "是否能追溯到条款、案例、公式与输入", "报告引用、计算明细、规则版本"],
        ["估算可用性", "汇总投资是否满足方案阶段快速研判", "同口径投资对比表"],
        ["现场可用性", "普通技术员是否能在弱网/断网条件完成流程", "操作录像/记录、任务耗时、问题单"],
        ["产品稳定性", "导入、计算、保存、恢复、导出是否稳定", "功能/异常/回归测试报告"],
        ["可扩展性", "新增单元、规则、地区单价是否需要改核心代码", "规则包扩展演示与接口测试"],
    ], widths=[2.5, 8.0, 5.5])

    doc.add_heading("十一、现场应用价值", level=1)
    add_table(doc, ["现场问题", "当前痛点", "系统作用", "预期效果"], [
        ["方案编制和调整慢", "查规范、找案例、等设计院，变更传递周期长", "图斑选取后按规则分钟级生成初步建议", "缩短内部研判周期，为正式设计提供结构化输入"],
        ["施工期措施漏项/错项", "一线人员非水保专业，工程、植物、临时措施易割裂", "按单元自动配齐措施并提示必要前置条件", "降低原则性漏项和返工风险"],
        ["投资估算依据不透明", "总价依赖经验，口径和价格基期不一致", "工程量、基价、调整系数和分类汇总全留痕", "快速支撑可研、变更和方案比选"],
        ["验收前突击补资料", "措施与验收证据未在施工期对应", "随措施自动生成核查项、责任和证据要求", "把验收要求前移到日常自查"],
        ["空间成果与台账割裂", "图斑、表格、照片、报告难关联", "依托 QGIS 工程把单元、决策、截图和报告关联", "一处定位、成套归档、可复核"],
        ["基层专业能力不足", "现场依赖少数专家", "固化专家规则并对低确定性工况主动升级复核", "扩大经验复用范围但不削弱专业责任"],
    ], widths=[2.5, 4.4, 5.2, 3.9])

    doc.add_heading("十二、成果体系与推广转化", level=1)
    doc.add_heading("（一）核心交付成果", level=2)
    add_table(doc, ["成果", "形态", "交付要求", "转化路径"], [
        ["水保措施智能辅助决策系统 V1.0", "C 端软件模块 + 离线部署包", "完成 P0 功能、测试、版本说明和安装手册", "抽蓄/输变电建设项目复制；后续作为课题五决策内核"],
        ["决策规则库 V1.0", "SQLite + JSON 规则包", "不少于 60 条规则/约束，来源与版本完整", "企业数据资产；可独立授权或持续维护"],
        ["综合单价库 V1.0", "数据库 + Excel 导入模板", "不少于 80 项，支持地区/基期/系数", "年度更新服务和区域扩展"],
        ["验收指标库 V1.0", "数据库 + 输出模板", "覆盖 V1.0 措施目录，关联六项指标", "施工自查、监理核查与验收准备"],
        ["报告模板包", "DOCX/XLSX/JSON", "建议书、估算表、核查表均可一键生成", "公司模板标准化和项目交付"],
        ["试点应用报告", "技术报告", "含方法、样本、对比、问题与修订证据", "产品推广案例和成果鉴定支撑"],
        ["知识产权与标准材料", "软著/论文/企业标准草案", "软著至少 1 项；论文 1 篇、标准草案 1 项为预期成果", "成果评价、报奖和市场推广"],
    ], widths=[3.0, 3.1, 5.1, 4.8])
    doc.add_heading("（二）产品版本路径", level=2)
    add_table(doc, ["阶段", "产品能力", "商业/推广形态"], [
        ["V1.0 专项预研", "4 类闭环、三库、离线 C 端、三件套报告、单站验证", "内部试用、软著、示范站点"],
        ["V1.5 区域推广", "扩展 8～10 类、规则管理工具、多区域价格包、更多案例验证", "项目授权 + 年度数据维护"],
        ["V2.0 课题五融合", "知识图谱/RAG、大模型对话填参、工具调用、服务 API", "企业级合规问答与决策平台"],
        ["V3.0 数字孪生融合", "施工进度、监测数据、无人机变化、整改闭环联动", "抽蓄建设期环水保数字孪生模块"],
    ], widths=[2.8, 8.2, 5.0])

    doc.add_heading("十三、实施计划与组织保障", level=1)
    doc.add_heading("（一）12 个月进度安排", level=2)
    add_table(doc, ["阶段", "月份", "主要工作", "里程碑/交付"], [
        ["阶段 1：需求与知识建模", "第 1～2 月", "现有 C 端适配设计；标准/案例清单；专家访谈；数据字典和分类框架", "M1：需求规格、数据模型、知识源清单"],
        ["阶段 2：三库与原型", "第 3～4 月", "规则建模、措施目录、单价/验收库；规则引擎和扰动单元原型", "M2：规则包 Alpha、核心引擎测试"],
        ["阶段 3：计算与产品集成", "第 5～7 月", "GIS 提参、参数计算、估算、验收映射、工作台和报告组件", "M3：可运行 Beta，完成 4 类主流程"],
        ["阶段 4：试点与评审", "第 8～10 月", "现场数据、3～5 单元运行、专家盲评、偏差分析和迭代", "M4：试点报告初稿、候选发布版"],
        ["阶段 5：定版与转化", "第 11～12 月", "回归测试、V1.0 冻结、部署/使用材料、软著和推广材料", "M5：V1.0 正式交付"],
    ], widths=[3.0, 2.2, 6.7, 4.1])

    doc.add_heading("（二）组织分工", level=2)
    add_table(doc, ["角色", "主要职责", "关键产出"], [
        ["项目负责人/产品负责人", "目标、范围、资源、里程碑和试点协调；最终验收", "任务书、决策记录、发布确认"],
        ["水保专业组", "标准/案例解析、规则编制、公式和验收指标、专家组织", "三库内容、规则说明、专业测试"],
        ["软件研发组", "模块设计、规则引擎、计算、GIS 集成、报告和安装包", "源代码、软件、测试和部署文档"],
        ["数据与测试组", "数据清洗、规则校验、测试用例、缺陷和版本管理", "数据包、测试报告、追溯清单"],
        ["试点单位", "提供脱敏资料、现场条件、操作反馈和对比依据", "试点数据、核查记录、应用证明"],
        ["专家组", "规则评审、盲评、重大差异裁决和成果咨询", "评审意见、评分表、签字记录"],
    ], widths=[3.1, 8.2, 4.7])

    doc.add_heading("（三）质量保障", level=2)
    add_bullet(doc, "需求基线：以 P0/P1 功能、输入输出和责任边界作为范围控制依据。")
    add_bullet(doc, "规则评审：专业编制、交叉校核、专家确认、发布审批四步闭环。")
    add_bullet(doc, "软件测试：核心规则和计算采用单元测试；关键流程采用回归、异常、断网和数据迁移测试。")
    add_bullet(doc, "数据质量：规则、单价和验收库均执行必填、唯一、单位、版本、来源和引用完整性校验。")
    add_bullet(doc, "发布管理：软件、三库和模板分别编号但在发布清单中绑定，确保试点结果可复现。")

    doc.add_heading("十四、经费概算", level=1)
    add_para(doc, "经费总额控制在 10 万元左右。考虑现有 QGIS/C++ 空间底座、ONNX 推理、智能分割和工程管理能力可直接复用，本次软件费用聚焦决策域与报告闭环；相应提高数据建库投入，确保形成真正可转化的规则资产。")
    add_table(doc, ["科目", "金额（万元）", "占比", "说明"], [
        ["软件开发与集成费", "4.0", "40%", "扰动单元、规则引擎、计算/估算、验收生成、C 端工作台、报告适配和测试"],
        ["数据建库费", "2.5", "25%", "标准/案例数字化，规则库、单价库、验收指标库和测试样本整理"],
        ["专家咨询费", "1.5", "15%", "3～5 名专家访谈、规则评审及 2～3 名专家试点盲评"],
        ["试点差旅与测试费", "1.0", "10%", "试点现场调研、资料采集、核查和测试环境"],
        ["软著及资料费", "0.5", "5%", "软著申报、报告/手册与成果材料"],
        ["不可预见费", "0.5", "5%", "价格数据补充、兼容性或现场测试调整"],
        ["合计", "10.0", "100%", "—"],
    ], widths=[3.1, 2.4, 1.8, 8.7])
    add_callout(doc, "经费使用原则", "预算不承担新建 GIS 底座、通用 Web 门户或大模型训练；若必须新增复杂结构计算、全国价格库或大规模多站验证，应另行追加经费并调整任务边界。", color=COLOR_AMBER, fill=COLOR_AMBER_LIGHT)

    doc.add_heading("十五、风险分析与对策", level=1)
    add_table(doc, ["风险", "可能性/影响", "主要对策", "触发后的范围调整"], [
        ["案例资料难获取或不可出域", "中/高", "提前签署数据清单；现场脱敏；只抽取规则特征，不复制敏感原文", "以不少于 5 个可用案例和专家补证替代"],
        ["试点工程进度不匹配", "中/高", "第 2 月前确定主备试点；允许已完成单元回溯验证", "切换备选站点或采用历史单元+现场复核"],
        ["规则数量膨胀导致延期", "高/中", "锁定 4 类 P0 闭环；其余仅做模型和模板", "P1 内容移入 V1.5，不挤占核心闭环"],
        ["单价口径无法统一", "高/中", "定义基期、地区、费率和材料价边界；同口径比较", "不评价单项绝对价格，只评汇总和计算链"],
        ["现有客户端编译/耦合风险", "中/中", "独立模块与核心库；优先复用稳定 QGIS API；设置回归基线", "必要时以独立动态库/进程适配，保持 DTO 不变"],
        ["报告被误作正式设计", "低/高", "报告自动附适用范围、复核等级和免责声明", "高风险工况只输出资料和专业复核要求"],
        ["标准更新", "中/中", "标准版本、有效期与规则包分离；发布前复核现行有效性", "通过新规则包升级，不修改历史结果"],
    ], widths=[3.1, 2.4, 6.8, 3.7])

    doc.add_heading("十六、与后续课题的衔接", level=1)
    add_para(doc, "本预研形成的规则库、单价库、验收指标库和计算服务是课题五可直接复用的决策知识底座；扰动单元、决策结果和报告 DTO 是后续系统接口契约。课题五完整项目可在此基础上叠加法规/标准文档 RAG、知识图谱和大模型对话层：大模型负责理解问题、补齐参数和解释结果，规则引擎负责确定性计算与合规约束，两者职责清晰。")
    add_para(doc, "与数字孪生平台衔接时，可将扰动单元和措施对象挂接到施工进度、无人机监测和整改任务，实现“识别变化—触发决策—安排措施—过程核查—指标评价”的时空闭环。由于核心逻辑与界面解耦，后续服务化不会推翻 V1.0 数据和规则资产。")
    add_table(doc, ["本预研资产", "课题五大模型系统", "数字孪生/其他平台"], [
        ["规则与计算服务", "作为工具调用，提供确定性建议和计算", "作为决策微服务或本地 SDK"],
        ["标准/案例来源元数据", "作为 RAG 索引和答案引用", "作为合规审计依据"],
        ["扰动单元与结果 DTO", "支持自然语言补参和结果解释", "关联时空对象、进度和监测"],
        ["单价/验收数据", "支持对话式估算和核查问答", "支持年度更新、任务派发和验收闭环"],
        ["试点证据与测试集", "用于评测大模型工具调用正确性", "用于跨站点复制验证"],
    ], widths=[4.1, 6.0, 5.9])
    add_callout(doc, "结论", "依托现有 C 端建设 V1.0，可以用有限经费把已具备的“扰动识别与核查”能力延伸为“识别—决策—估算—验收—报告”产品闭环。项目的首要成果不是界面数量，而是一套经试点验证、可追溯、可持续迭代并能被后续平台调用的水土保持决策内核。")

    doc.add_page_break()
    doc.add_heading("附录 A：V1.0 功能优先级清单", level=1)
    add_table(doc, ["优先级", "功能", "V1.0 交付判定"], [
        ["P0", "抽蓄水保工程模板", "可新建工程并管理项目信息、目录、图层与决策数据"],
        ["P0", "扰动单元台账", "识别/绘制图斑可转单元；支持状态、参数、版本和地图定位"],
        ["P0", "GIS 参数提取", "面积、周长、基础坡度/高差等可自动提取并记录来源"],
        ["P0", "引导式参数录入", "按 4 类核心单元显示条件字段、单位、缺项和校验"],
        ["P0", "规则引擎", "支持条件、优先级、依赖/互斥、版本、命中链和复核门"],
        ["P0", "关键参数计算", "水力、典型工程量、边坡/植物推荐按 V1.0 规则运行"],
        ["P0", "投资估算", "单价套用、调整系数、费用分类、方案对比和 Excel 输出"],
        ["P0", "验收清单", "措施映射、核查方式/依据、责任与证据字段、Excel 输出"],
        ["P0", "Word 建议书", "包含单元概况、措施、计算、投资、验收、依据和限制"],
        ["P0", "保存与追溯", "保存输入快照、规则/价格/模板版本、人工修改和输出文件"],
        ["P0", "离线部署", "无网络完成核心功能，规则包可离线导入并校验"],
        ["P1", "8～10 类全部深度规则", "完成数据模型和扩展模板，深度内容可进入 V1.5"],
        ["P1", "可视化规则编辑器", "V1.0 仅需受控导入、校验和版本查看"],
        ["P1", "多站协同与用户权限", "V1.0 单机项目级使用，企业协同后续建设"],
        ["P1", "知识图谱/RAG/大模型问答", "仅保留数据与工具接口，不在本次经费内"],
        ["P1", "移动端/完整 B/S 系统", "不在 V1.0 建设范围"],
    ], widths=[2.0, 5.2, 8.8])

    doc.add_heading("附录 B：规则数据结构示例", level=1)
    add_para(doc, "以下为逻辑结构示例，用于说明规则资产必须包含条件、输出、来源和复核门。实际交付采用经过评审的数据字典与 JSON Schema。", indent=False)
    code = """{
  \"rule_id\": \"SP-GULLY-001\",
  \"version\": \"1.0.0\",
  \"unit_type\": \"渣场\",
  \"subtype\": \"沟道型\",
  \"conditions\": [
    {\"field\": \"spoil_volume_m3\", \"op\": \">\", \"value\": 500000}
  ],
  \"outputs\": [
    {\"measure\": \"拦挡工程\", \"required\": true},
    {\"measure\": \"截排水系统\", \"required\": true},
    {\"measure\": \"堆体整形与植被恢复\", \"required\": true}
  ],
  \"review_gate\": \"高风险工程须专项设计及稳定复核\",
  \"sources\": [{\"document\": \"现行标准/案例/专家纪要\", \"clause\": \"待绑定\"}],
  \"status\": \"reviewed\"
}"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F5F6")
    set_cell_border(cell, COLOR_LINE, 6)
    set_cell_margins(cell, 180, 220, 180, 220)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(code.splitlines()):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(COLOR_DARK)
    add_para(doc, "说明：本示例不固化具体工程设计结论。“待绑定”内容在规则发布前必须替换为现行有效依据、条款或经签字确认的专家纪要；计算类输出还需关联公式编号、单位和适用范围。", color=COLOR_GRAY, size=9.5, indent=False)

    doc.add_heading("附录 C：主要依据与版本管理原则", level=1)
    add_para(doc, "V1.0 知识源应至少覆盖下列类别。由于标准、定额和管理文件会更新，项目启动和 V1.0 发布前应由专业组再次核验名称、编号、现行有效状态、适用范围及替代关系，系统规则不得仅凭文件名长期固化。", indent=False)
    add_table(doc, ["类别", "建议纳入的主要依据（以现行有效版本为准）", "数字化重点"], [
        ["生产建设项目技术", "《生产建设项目水土保持技术标准》GB 50433 等", "防治责任范围、措施总体布设、设计要求"],
        ["防治目标", "《生产建设项目水土流失防治标准》GB/T 50434 等", "目标等级、六项防治指标及地区修正"],
        ["工程设计", "《水土保持工程设计规范》GB 51018 等", "工程措施选型、参数、设计频率和校核要求"],
        ["水利水电工程", "《水利水电工程水土保持技术规范》SL 575 等", "抽蓄/水利水电施工区、渣场、道路等行业要求"],
        ["监测评价", "《生产建设项目水土保持监测与评价标准》GB/T 51240 等", "监测对象、方法、评价和数据留痕"],
        ["质量与验收", "水土保持工程质量评定、生产建设项目水土保持设施自主验收相关现行规程和文件", "单位/分部/单元工程、核查方法、资料和判定依据"],
        ["概（估）算", "水土保持工程概（估）算编制规定、定额及项目所在地现行计价文件", "项目划分、单价构成、费率、地区和基期调整"],
        ["企业与流域要求", "水利部、流域管理机构、国网公司及建设单位现行环水保制度", "企业控制要求、审批/复核节点和表单"],
        ["工程案例", "5～10 个抽蓄电站方案、变更、监理、监测和验收材料", "条件—措施—工程量—问题—整改的结构化实例"],
        ["专家经验", "3～5 名专家访谈和评审纪要", "长尾工况、优先级、禁配关系、复核门"],
    ], widths=[2.8, 8.5, 4.7])
    add_para(doc, "版本管理原则：规则包、单价包、验收包和模板包分别版本化；每次决策绑定完整版本组合。标准更新时发布新规则版本，历史结果保持原版本可追溯；跨版本比较必须显示规则和价格变化。", bold_prefix="版本管理原则：", indent=False)

    return doc


def validate_document(path: Path) -> None:
    if not path.exists() or path.stat().st_size < 50_000:
        raise RuntimeError("DOCX output is missing or unexpectedly small")
    with ZipFile(path) as archive:
        required = {"[Content_Types].xml", "word/document.xml", "word/styles.xml"}
        missing = required.difference(archive.namelist())
        if missing:
            raise RuntimeError(f"Invalid DOCX, missing: {sorted(missing)}")
    opened = Document(path)
    headings = [p.text for p in opened.paragraphs if p.style.name.startswith("Heading")]
    required_headings = ["一、项目摘要", "六、研究任务与主要内容", "十四、经费概算", "附录 C：主要依据与版本管理原则"]
    for heading in required_headings:
        if heading not in headings:
            raise RuntimeError(f"Required heading missing: {heading}")
    table_text = [cell.text for table in opened.tables for row in table.rows for cell in row.cells]
    full_text = "\n".join([*(p.text for p in opened.paragraphs), *table_text])
    for phrase in ("QGIS/C++", "10 万元", "规则引擎", "投资估算", "验收核查表"):
        if phrase not in full_text:
            raise RuntimeError(f"Required phrase missing: {phrase}")
    if len(opened.tables) < 20:
        raise RuntimeError("Expected proposal tables are missing")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    create_architecture_diagram(ASSET_DIR / "target_architecture.png")
    create_flow_diagram(ASSET_DIR / "technical_flow.png")
    doc = build_document()
    doc.save(OUTPUT_FILE)
    validate_document(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
